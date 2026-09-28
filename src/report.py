"""Build the DOCX report (reports/Freight_Rate_Report.docx) with python-docx.

Numbers come from reports/metrics/*.csv|json (written by the pipeline), figures from
reports/figures/ and scorer_results/candidate_december.png.
Run: python -m src.report
"""
import json

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .config import FIGURES, METRICS, REPORTS, ROOT

OUT = REPORTS / "Freight_Rate_Report.docx"
ACCENT = RGBColor(0x06, 0x4A, 0x56)
FINAL_ID = "A34"

# key-model rows shown in the report, with plain-language labels
ROWS = [
    ("B1 global median rpm x distance", "Baseline: global median $/mile x distance"),
    ("B2 lane median rpm", "Baseline: lane median $/mile"),
    ("B3 ridge (log rpm)", "Baseline: ridge on log $/mile"),
    ("M0 LightGBM log-rpm", "LightGBM, all features (first model)"),
    ("A24 cycle - holidays + additive calendar", "Two-stage: tree + weekday/quarter-end model"),
    ("A40 A39 + recency half-life 60d", "+ coarse quote regime + recency (smooth alt.)"),
    ("A42 A24 + qs_7d + quote x regime", "+ trailing quote x regime (smooth alt.)"),
    ("A34 A27 + recency half-life 60d", "FINAL: + daily quote regime + recency weights"),
]


# ------------------------------------------------------------------ helpers
def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")  # CLEAR, never SOLID
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def table(doc, header, rows, widths, bold_last=False, font=8.5):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(font), RGBColor(0xFF, 0xFF, 0xFF)
        shade(c, "064A56")
    for k, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(font)
            run.bold = bold_last and k == len(rows) - 1
            if i > 0 and str(v).replace(",", "").replace(".", "").replace("-", "").isdigit():
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT  # numbers only
            if bold_last and k == len(rows) - 1:
                shade(cells[i], "E6F0F1")
    for row in t.rows:  # widths on every cell (Word honours cell widths, not column widths)
        for i, w in enumerate(widths):
            row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            p.add_run(it[0]).bold = True
            p.add_run(it[1])
        else:
            p.add_run(it)
        p.paragraph_format.space_after = Pt(1)


def para(doc, text, italic=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(4)
    return p


def figure(doc, path, caption, width=6.5):
    doc.add_picture(str(path), width=Inches(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.keep_with_next = True  # keep the caption on the same page
    para(doc, caption, italic=True, size=8.5).alignment = WD_ALIGN_PARAGRAPH.CENTER


def money(x):
    return f"{x:,.1f}"


# ------------------------------------------------------------------ document
def build():
    key = pd.read_csv(METRICS / "key_models.csv", index_col=0)
    clean_log = pd.read_csv(METRICS / "cleaning_log.csv")
    ens = pd.read_csv(METRICS / "ensemble_check.csv")
    dec = json.loads((METRICS / "december_summary.json").read_text())
    final = key.loc[[k for k, _ in ROWS if k.startswith(FINAL_ID)][0]]
    a24 = key.loc["A24 cycle - holidays + additive calendar"]
    m0 = key.loc["M0 LightGBM log-rpm"]
    ridge = key.loc["B3 ridge (log rpm)"]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.9))
    sec.top_margin = sec.bottom_margin = Inches(0.8)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(10)
    for name, size in (("Title", 20), ("Heading 1", 13), ("Heading 2", 11)):
        doc.styles[name].font.color.rgb = ACCENT
        doc.styles[name].font.size = Pt(size)

    doc.add_heading("Freight Load Rate Prediction", level=0)
    para(doc, "Spotter - Machine Learning Engineer assessment. Code, metrics and figures are reproducible with "
              "`python -m src.pipeline` (see README).", italic=True, size=9)

    # ---- summary
    doc.add_heading("Summary", level=1)
    bullets(doc, [
        ("Task. ", "Predict posted_rate for 12,000 loads in Nov-Dec 2025 from 48,000 labelled loads in Jan-Oct 2025, "
                   "and a daily December curve for one fixed lane."),
        ("Validation. ", "Rolling-origin time folds with a 2-month horizon (the real setup), plus a city-holdout "
                         "and a random split for comparison. Every learned statistic is refitted inside each fold."),
        ("Model. ", "Two stages on log rate-per-mile: a LightGBM tree for the load-level price (lane/city/region "
                    "encodings, distance, equipment, weight, quote_signal and its daily regime) plus a small ridge "
                    "model for day-level effects (weekday, quarter-end ramp)."),
        ("Result. ", f"Time-fold MAE ${final['time mean']:.1f} (MAPE {final['time MAPE %']:.2f}%), worst fold "
                     f"${final['time worst']:.1f}, vs ${ridge['time mean']:.1f} for ridge and ${m0['time mean']:.1f} "
                     "for a plain LightGBM on all features."),
        ("Key finding. ", "quote_signal looked useless (pooled correlation ~0) but its relation to price flips "
                          "sign between regimes that the daily mean quote identifies; modelling the regime is the "
                          "largest single accuracy gain."),
    ])

    # ---- data
    doc.add_heading("1. Data and key findings", level=1)
    bullets(doc, [
        ("Split in time. ", "Training covers 2025-01-01..10-31, validation 11-01..12-31: the task is a 2-month-ahead "
                            "forecast. 8 cities appear only in validation (1,447 loads, 12.1%); every other validation "
                            "load is on a lane seen in training (4,014 lanes, median 10 loads each)."),
        ("Price structure. ", "Rate per mile falls with distance (~$2.65 below 330 mi, ~$1.89 above 2,200 mi); "
                              "Reefer > Flatbed > Dry Van; heavier loads pay up to ~6% more."),
        ("market_index ", "is one value per day plus small per-load noise (within-day sd 0.025 vs 0.168 overall), with "
                          "a strong weekly cycle (Thu high, Sun low). Its slow level is confounded with time: low in "
                          "Jan-Feb when prices were low, low again in Aug-Oct when they were not."),
        ("Calendar. ", "Prices rise over the last ~3 weeks of each quarter (Mar, Jun, Sep) and drop after; weekday "
                       "effect about +/-0.8%; a ~5% upward drift Jan-Oct not explained by the index. No visible "
                       "public-holiday effect."),
        ("quote_signal regimes. ", "Pooled correlation with price is -0.04, so it first looked like noise, not "
                                   "leakage. Per week, on long loads, the correlation ranges -0.32..+0.49 and tracks "
                                   "the weekly mean quote (r = 0.87): positive when the mean quote is high (> 2.1: late "
                                   "Feb, Mar, Jun, Sep), negative when low (< 2.0: Apr, May, Jul, Oct), ~0 in between "
                                   "(Aug). The regimes cancel out when pooled. Nov-Dec sit at ~2.05, the transition zone."),
    ])
    figure(doc, FIGURES / "12_quote_regime.png",
           "Figure 1. Left: weekly mean quote_signal forms three regimes. Right: the quote-price correlation per week "
           "follows the regime.")

    # ---- data quality
    doc.add_heading("2. Data-quality issues and fixes", level=1)
    para(doc, "Validation rows are never dropped; feature problems are fixed or imputed and flagged. Only training rows "
              "with a corrupted target are removed. Fitted rules (weight medians, the corrupted-label filter) are "
              "refitted inside every CV fold.")
    rows = [(r.rule, r.action, f"{r.train_rows:,}", f"{r.valid_rows:,}") for r in clean_log.itertuples()]
    table(doc, ["Rule", "Fix", "Train rows", "Valid rows"], rows, [1.9, 3.2, 0.8, 0.8], font=8)
    bullets(doc, [
        ("Corrupted rates. ", "677 training rates (1.4%) are x2-x6 away from the expected price, spread evenly over "
                              "equipment, month, weekday, distance and city: random multiplicative errors. They are "
                              "flagged by an out-of-fold robust (L1) GBM that uses no lane identity; a lane-median rule "
                              "agreed on 673 rows but mis-flagged thin lanes (1-2 loads) where the bad row itself moves "
                              "the median. Sanity check: Huber loss on uncleaned labels matches L2 on cleaned labels, "
                              "plain L2 on uncleaned labels is much worse."),
        ("Coordinates ", "are distorted and partly clipped (e.g. Boston/Providence lon = -69.5); distance is consistent "
                         "(~1.18x great-circle, +/-5% within lane), so distance is authoritative."),
        ("No duplicates, ", "unparseable dates or spelling variants were found (checked defensively)."),
    ])

    # ---- validation
    doc.add_heading("3. Validation approach", level=1)
    bullets(doc, [
        ("Time folds (primary). ", "T1 Jan-Apr > May-Jun, T2 Jan-Jun > Jul-Aug, T3 Jan-Aug > Sep-Oct: a 2-month horizon "
                                   "after the training cut-off, mirroring Jan-Oct > Nov-Dec. The model is chosen on the "
                                   "worst fold, not only the mean."),
        ("City holdout. ", "64 cities in 8 random groups; each fold holds out every load touching a group city as "
                           "pickup OR delivery. It measures the unseen-city risk (12% of validation)."),
        ("Random 5-fold ", "only as a contrast: train and test share dates, so it rewards memorising day-level prices; "
                           f"its MAE is {m0['time mean'] / m0['random MAE']:.1f}x lower than the time folds for the first "
                           f"LightGBM and {final['time mean'] / final['random MAE']:.1f}x lower for the final model."),
        ("Leakage guard. ", "Weight medians, the corrupted-label filter, target encodings and their priors, the "
                            "unseen-city blanking (training rows only), the stage-2 model and the tree are refitted in "
                            "every fold on that fold's training rows. Test rows get encodings from the training period only."),
        ("Labels. ", "Metrics are reported on clean labels and on raw labels (corrupted rows kept), since the scored "
                     "labels likely contain the same ~1.4% corruption."),
    ])

    # ---- model
    doc.add_heading("4. Model", level=1)
    bullets(doc, [
        ("Target: ", "log(rate per mile); predictions are exp(.) x distance (log vs rpm vs $ targets were compared; "
                     "smearing changed nothing measurable)."),
        ("Stage 1 - LightGBM: ", "distance, great-circle distance, coordinates, equipment, weight (+ cap/floor/missing "
                                 "flags), per-load market-index deviation, quote_signal and the day's mean quote "
                                 "(regime), and out-of-fold smoothed target encodings of lane, pickup, delivery, 5-degree "
                                 "region and region pair. Unseen keys fall back to an equipment x distance-band prior; "
                                 "12% of training rows get one city blanked so the model learns that case. Recency "
                                 "weights with a 60-day half-life."),
        ("Stage 2 - day effects: ", "ridge on the tree's date-grouped out-of-fold residuals: weekday dummies + a linear "
                                    "ramp over the last 21 days of a quarter. Keeping date features out of the tree "
                                    "stops it from memorising day-level prices that do not carry over to new dates."),
        ("Tested and rejected: ", "slow market-index level (-3..-5% bias two months out), holiday features (no gain; "
                                  "Thanksgiving/Christmas never seen), linear time trend (unstable: much worse on T2), "
                                  "level offsets from recent residuals (overshoot after the June peak), a ridge blend "
                                  f"(weight tuned on T1-T2 = {ens['tuned_weight'].iloc[0]:.2f}; failed on T3: "
                                  f"${ens.iloc[2]['T3: Jan-Aug > Sep-Oct']:.1f} vs ${ens.iloc[0]['T3: Jan-Aug > Sep-Oct']:.1f})."),
    ])

    # ---- results
    doc.add_heading("5. Results", level=1)
    rows = []
    for k, label in ROWS:
        r = key.loc[k]
        rows.append((label, money(r["T1"]), money(r["T2"]), money(r["T3"]), money(r["time mean"]),
                     money(r["time worst"]), f"{r['time MAPE %']:.2f}", money(r["city MAE"]), money(r["random MAE"])))
    table(doc, ["Model", "T1", "T2", "T3", "Mean", "Worst", "MAPE %", "City", "Random"], rows,
          [2.45, 0.5, 0.5, 0.5, 0.52, 0.52, 0.55, 0.5, 0.56], bold_last=True, font=8)
    para(doc, "MAE in $ on clean labels. T1-T3 = time folds; City = city-holdout mean; Random = random 5-fold mean.",
         italic=True, size=8.5)
    bullets(doc, [
        ("Versus baselines: ", f"mean time-fold MAE is {100 * (1 - final['time mean'] / ridge['time mean']):.0f}% below "
                               f"ridge and {100 * (1 - final['time mean'] / key.loc['B2 lane median rpm', 'time mean']):.0f}% "
                               "below the lane median, and lower than the first LightGBM on every fold. Ridge remains "
                               f"better on T2 (${ridge['T2']:.1f} vs ${final['T2']:.1f}) but fails on T3 "
                               f"(${ridge['T3']:.1f}, -5% bias when the market index is low)."),
        ("Raw labels: ", f"time-fold MAE ${final['time MAE raw']:.1f} and RMSE ${final['time RMSE raw']:.0f} with "
                         f"corrupted rows kept (clean: ${final['time mean']:.1f} / ${final['time RMSE clean']:.0f}). "
                         "RMSE on raw labels is dominated by the ~1.4% corrupted rows for every model."),
        ("Unseen cities: ", f"city-holdout MAE ${final['city MAE']:.1f} vs ${final['random MAE']:.1f} on the random "
                            "split: the penalty for an unknown city is small (~$3)."),
        ("Where it errs: ", "the weakest fold is T2 (Aug under-predicted ~3%: August's quote regime was neutral while "
                            "January, the only similar training month, was positive). MAPE is highest on short loads."),
    ])

    # ---- December
    doc.add_heading("6. December chart", level=1)
    para(doc, "The December file has no coordinates, market_index or quote_signal. Coordinates are looked up per city "
              "(identical for every load of a city); market_index and the daily quote regime use the real daily means "
              f"of validation.csv (input features, not targets; a forecast from October missed the index by 0.02-0.07); "
              f"quote_signal is the median of comparable training loads (Dry Van, 300-420 mi: {dec['quote_fill']:.3f}). "
              "December 2025 lies outside the training range, so the model contains no absolute-time feature a tree "
              "would have to extrapolate.")
    figure(doc, ROOT / "scorer_results" / "candidate_december.png",
           "Figure 2. candidate_december.png produced by score.py.", width=6.3)
    s2 = dec["stage2_pct"]
    bullets(doc, [
        ("Range ", f"${dec['min']:.0f}-${dec['max']:.0f} (mean ${dec['mean']:.0f}; recent Aug-Oct Dry Van loads on this "
                   "lane scaled to 360 mi average ~$808)."),
        ("Stage-2 day effects ", f"are small in the final model (Thursday +{s2['dow_Thu']:.1f}% vs Monday, quarter-end "
                                 f"ramp +{s2['quarter_end_ramp']:.1f}% by 31 Dec): the daily quote mean also carries "
                                 "day-level information, so the tree absorbs much of that movement."),
        ("Day-to-day jitter: ", f"the load-level part moves between ${dec['tree_min']:.0f} and ${dec['tree_max']:.0f} "
                                f"because the daily mean quote ({dec['qs_day_min']:.3f}-{dec['qs_day_max']:.3f}) sits "
                                "on the regime boundary, where the tree switches between the positive (Jan-like) and "
                                "neutral (Aug-like) regime. Smoothed regime features (coarse bins, trailing 7/14-day "
                                "means, quote x regime interaction) reduced or removed the jitter but cost 8-21% mean "
                                "MAE on the time folds (a fully smooth weekday + quarter-end curve cost 21%); since "
                                "validation accuracy is what is scored, the more accurate model was kept, and the same "
                                "model produced both files."),
    ])

    # ---- limitations
    doc.add_heading("7. Limitations and next steps", level=1)
    bullets(doc, [
        ("No Christmas/year-end period in training: ", "Thanksgiving and Christmas are predicted like ordinary days "
                                                       "(e.g. 25 Dec as a Thursday)."),
        ("Quarter-end effect ", "is learned from only three quarter-ends (Mar, Jun, Sep); the year-end may differ."),
        ("Quote regime ambiguity: ", "Nov-Dec sit at the regime boundary (~2.05), where training has only Jan "
                                     "(positive) and Aug (neutral) as examples; this is the largest source of risk "
                                     "and of the December jitter."),
        ("Price-level drift ", "is only partly predictable; level errors of 1-3% per month remain the main error "
                               "source on the time folds."),
        ("Next steps: ", "a longer history (a previous December), an explicit regime model (e.g. hidden-Markov on the "
                         "daily quote), quantile outputs for pricing ranges, and monitoring of the regime indicator."),
    ])
    doc.save(OUT)
    return OUT


def main():
    path = build()
    print(f"Wrote {path}")


if __name__ == "__main__":
    main()
